#!/usr/bin/env python
"""Merge manifests, convert raw files to text, slice clauses, check params, build sqlite, write STATUS.md.

Python 3.9, conda env "pytorch".   python official/tools/build.py [--dense]   (ends by running graph.py)
doc / rtf / xls are converted with LibreOffice (headless); results are cached in official/_conv/.
"""
import csv
import hashlib
import io
import re
import shutil
import sqlite3
import subprocess
import sys
import zipfile
from pathlib import Path

import fitz                      # PyMuPDF
from bs4 import BeautifulSoup

sys.path.insert(0, str(Path(__file__).resolve().parent))
from xfa_text import is_xfa, xfa_text    # noqa: E402  dynamic (XFA) forms keep their text in an XML template

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
csv.field_size_limit(10 ** 9)
ROOT = Path(__file__).resolve().parents[1]
TEXT, CONV = ROOT / "text", ROOT / "_conv"
ENC = "utf-8-sig"
SOFFICE = r"C:\Program Files\LibreOffice\program\soffice.exe"

# Only these pages of a PDF are sliced into articles (the rest are annexes holding other instruments).
PAGE_RANGE = {"treaty.cn-sg.2007.en": (1, 20)}
# Pages kept in the text layer. Annex D of the IRAS file (pp. 28-46) is the terminated 1986 agreement.
PAGE_KEEP = {"treaty.cn-sg.2007.en": [(1, 27), (47, 47)]}

CN_NUM = {"零": 0, "〇": 0, "一": 1, "二": 2, "两": 2, "三": 3, "四": 4, "五": 5, "六": 6, "七": 7, "八": 8, "九": 9}


def cn2int(s):
    total, cur = 0, 0
    for ch in s:
        if ch in CN_NUM:
            cur = CN_NUM[ch]
        elif ch == "十":
            total += (cur or 1) * 10
            cur = 0
        elif ch == "百":
            total += (cur or 1) * 100
            cur = 0
        else:
            return None
    return total + cur


def read_csv(path):
    if not path.exists():
        return []
    with path.open(encoding=ENC, newline="") as f:
        return list(csv.DictReader(f))


def write_csv(path, fields, rows):
    with path.open("w", encoding=ENC, newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


# ---------------------------------------------------------------- 1. merge manifests
def merge_manifests():
    rows, seen, problems = [], {}, []
    verified = {v["id"]: v["match"] for v in read_csv(ROOT / "verify.csv")}
    for p in sorted(ROOT.glob("manifest_*.csv")):
        for r in read_csv(p):
            f = ROOT / r["file"]
            if not f.exists():
                problems.append("missing file: %s" % r["file"])
                continue
            if hashlib.sha256(f.read_bytes()).hexdigest() != r["sha256"]:
                problems.append("sha256 mismatch: %s" % r["id"])
            if r["via"] == "webfetch_copy" and (r["id"].endswith(".full") or r["id"].endswith(".data")):
                r["via"], r["verbatim"] = "assembled_local", "derived_from_official_raw"
            elif r["via"] == "webfetch_copy" and "RAW bytes via curl" in r["note"]:
                r["via"] = "curl_cli"
                r["verbatim"] = "raw" if verified.get(r["id"]) == "yes" else "raw_not_reverified"
            m = re.search(r"parent=([A-Za-z0-9._-]+)", r["note"])
            if m and not r["parent"]:
                r["parent"] = m.group(1)
            if r["id"] in seen:
                problems.append("duplicate id, later row kept: %s" % r["id"])
                rows[seen[r["id"]]] = r
                continue
            seen[r["id"]] = len(rows)
            rows.append(r)
    by_hash = {}
    for r in rows:
        by_hash.setdefault(r["sha256"], []).append(r["id"])
    alias = {}
    for ids in by_hash.values():
        for other in ids[1:]:
            alias[other] = ids[0]                      # identical bytes: slice only the first id
    for r in rows:
        r["alias_of"] = alias.get(r["id"], "")
    return rows, problems


# ---------------------------------------------------------------- 2. raw -> text
def real_kind(path, ext):
    head = path.read_bytes()[:8]
    ext = ext.lower()
    if head.startswith(b"%PDF-"):
        return "pdf"
    if head.startswith(b"{\\rtf"):
        return "rtf"
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        return "xls" if ext in ("xls", "et") else ("ppt" if ext in ("ppt", "pps") else "doc")
    if head.startswith(b"PK"):
        return "xlsx" if ext in ("xlsx", "xls") else ("pptx" if ext == "pptx" else "docx")
    return ext


def html_text(raw):
    soup = BeautifulSoup(raw, "lxml")
    for t in soup(["script", "style", "noscript", "header", "footer", "nav", "form", "iframe"]):
        t.decompose()
    body = soup.body or soup
    ps = body.find_all("p")
    total = sum(len(p.get_text(strip=True)) for p in ps)
    node = body
    if total > 400:                                   # descend to the smallest container holding most <p> text
        while True:
            best = None
            for child in node.find_all(recursive=False):
                if child.name == "p":
                    continue
                share = sum(len(p.get_text(strip=True)) for p in child.find_all("p"))
                if share >= 0.7 * total:
                    best = child
            if best is None:
                break
            node = best
    text = node.get_text("\n")
    lines = [re.sub(r"[ \t\u00a0\u3000]+", " ", ln).strip() for ln in text.splitlines()]
    return "\n".join(ln for ln in lines if ln)


def pdf_pages(path):
    with fitz.open(path) as doc:
        return [page.get_text("text") for page in doc]


def docx_text(path):
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode("utf-8", "replace")
    xml = re.sub(r"</w:p>", "\n", xml)
    return re.sub(r"<[^>]+>", "", xml)


def xlsx_text(path):
    import openpyxl
    wb = openpyxl.load_workbook(str(path), data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append("## sheet: %s" % ws.title)
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                out.append(" | ".join(cells))
    return "\n".join(out)


def soffice_convert(jobs):
    """jobs: list of (src_path, kind, sha16). Writer formats -> txt, calc formats -> xlsx. Cached in _conv/."""
    CONV.mkdir(exist_ok=True)
    tmp = CONV / "in"
    for target, kinds in (("txt:Text (encoded):UTF8", ("doc", "rtf")), ("xlsx", ("xls",)), ("pdf", ("ppt", "pptx"))):
        todo = []
        out_ext = {"xlsx": "xlsx", "pdf": "pdf"}.get(target, "txt")
        for src, kind, sha in jobs:
            if kind in kinds and not (CONV / ("%s.%s" % (sha, out_ext))).exists():
                tmp.mkdir(exist_ok=True)
                dst = tmp / ("%s.%s" % (sha, kind))
                shutil.copyfile(str(src), str(dst))
                todo.append(str(dst))
        for i in range(0, len(todo), 8):
            profile = "file:///" + str(CONV / "lo_profile").replace("\\", "/")
            subprocess.run([SOFFICE, "--headless", "--norestore", "-env:UserInstallation=" + profile,
                            "--convert-to", target, "--outdir", str(CONV)] + todo[i:i + 8],
                           stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=1800)
    if tmp.exists():
        shutil.rmtree(str(tmp), ignore_errors=True)


def to_text(row, kind):
    """Returns (text, pages or None, status)."""
    path = ROOT / row["file"]
    sha = row["sha256"][:16]
    try:
        if kind in ("html", "htm"):
            return html_text(path.read_bytes()), None, "ok"
        if kind == "pdf":
            pages = pdf_pages(path)
            if row["id"] in PAGE_KEEP:
                pages = [t for i, t in enumerate(pages, 1) if any(a <= i <= b for a, b in PAGE_KEEP[row["id"]])]
            text = "\n".join(pages)
            if len(text) < 3000 and ("Please wait" in text or is_xfa(path)):
                t = xfa_text(path)
                return (t, None, "xfa_template") if t else (text, pages, "xfa_form_no_text")
            return text, pages, "ok" if len(text.strip()) > 200 else "scanned_or_empty"
        if kind == "docx":
            return docx_text(path), None, "ok"
        if kind == "xlsx":
            return xlsx_text(path), None, "ok"
        if kind in ("doc", "rtf"):
            f = CONV / (sha + ".txt")
            return (f.read_text(encoding="utf-8-sig", errors="replace"), None, "ok") if f.exists() \
                else ("", None, "convert_failed")
        if kind == "xls":
            f = CONV / (sha + ".xlsx")
            return (xlsx_text(f), None, "ok") if f.exists() else ("", None, "convert_failed")
        if kind in ("ppt", "pptx"):
            f = CONV / (sha + ".pdf")
            if not f.exists():
                return "", None, "convert_failed"
            pages = pdf_pages(f)
            return "\n".join(pages), pages, "ok"
        if kind in ("md", "txt", "json"):
            return path.read_text(encoding="utf-8", errors="replace"), None, "ok"
        return "", None, "unsupported:" + kind
    except Exception as e:                            # noqa: BLE001
        return "", None, "error:%s" % type(e).__name__


# ---------------------------------------------------------------- 3. slice clauses
RE_CN_ART = re.compile(r"^第([一二三四五六七八九十百零〇两]+)[条條]\s*(.*)$")
RE_CN_ITEM = re.compile(r"^([一二三四五六七八九十]+)、\s*(.*)$")
RE_EN_ART = re.compile(r"^ARTICLE\s+(\d+)\b\s*(.*)$", re.I)
RE_EN_PARA = re.compile(r"^(\d{1,2})\.(?:\s+(\S.*))?$")      # "1." may sit alone on a line in PDF text
RE_PAGENO = re.compile(r"^-?\s*\d{1,3}\s*-?$")


RE_CHAPTER = re.compile(r"^第[一二三四五六七八九十百]+[章节節]\s*\S{0,20}(\s\S{0,6}){0,6}$")
RE_SPLIT_HEAD = re.compile(r"^第[一二三四五六七八九十百零〇两]*$")
RE_SPLIT_TAIL = re.compile(r"^[一二三四五六七八九十百零〇两]*[条條]")


def clean_lines(text):
    lines = [ln.strip() for ln in text.splitlines() if ln.strip() and not RE_PAGENO.match(ln.strip())]
    lines = [ln for ln in lines if not RE_CHAPTER.match(ln)]
    out, i = [], 0
    while i < len(lines):                              # "第" + "二十三条 ..." were split by an inline tag
        if i + 1 < len(lines) and RE_SPLIT_HEAD.match(lines[i]) and RE_SPLIT_TAIL.match(lines[i + 1]):
            out.append(lines[i] + lines[i + 1])
            i += 2
        else:
            out.append(lines[i])
            i += 1
    return out


def split_paras(block, matcher, number):
    """Split an article body into consecutively numbered paragraphs."""
    paras, cur, buf = [], None, []
    for ln in block:
        m = matcher.match(ln)
        if m and number(m.group(1)) == (cur or 0) + 1:
            if cur is not None:
                paras.append((cur, buf))
            cur, buf = number(m.group(1)), [ln]
        elif cur is not None:
            buf.append(ln)
    if cur is not None:
        paras.append((cur, buf))
    return paras


def pick_run(lines, regex, number):
    """Headings must be numbered consecutively; a jump is accepted only for a short title-like line.
    A restart at 1 opens a new run (table of contents first, or a second instrument in an annex).
    Returns the first run whose blocks are real text, as [(line_index, number)], plus the end line."""
    runs, cur, last = [], [], 0
    for i, ln in enumerate(lines):
        m = regex.match(ln)
        if not m:
            continue
        no = number(m.group(1))
        if no is None:
            continue
        title_like = len(ln) <= 24 and not re.search(r"[，。；：,;]", ln)
        if no == 1 and last >= 2:
            runs.append(cur)
            cur, last = [], 0
        if no == last + 1 or (no > last + 1 and title_like):
            cur.append((i, no))
            last = no
    runs.append(cur)
    runs = [r for r in runs if r]
    if not runs:
        return [], len(lines)
    bounds = [r[0][0] for r in runs] + [len(lines)]
    best = None
    for k, r in enumerate(runs):
        chars = sum(len(x) for x in lines[r[0][0]:bounds[k + 1]])
        if chars >= 150 * len(r) and len(r) >= 3:
            best = k
            break
    if best is None:
        best = max(range(len(runs)), key=lambda k: len(runs[k]))
    return runs[best], bounds[best + 1]


RE_PROTOCOL = re.compile(r"^.{0,30}(议\s*定\s*书|議\s*定\s*書)$|^PROTOCOL$", re.I)


def slice_articles(doc_id, text, regex, number, para_regex, para_number, fmt_art, fmt_para, mode):
    lines = clean_lines(text)
    run, end = pick_run(lines, regex, number)
    out = []
    if len(run) < 3:
        return out, "none"
    for k, (i, no) in enumerate(run):
        stop = run[k + 1][0] if k + 1 < len(run) else end
        block = lines[i:stop]
        if k + 1 == len(run):                          # a protocol may follow the last article
            for j, ln in enumerate(block[1:], 1):
                if RE_PROTOCOL.match(ln):
                    out.append({"cite_id": doc_id + "#protocol", "doc_id": doc_id, "kind": "protocol",
                                "label": "议定书", "parent": "", "text": "\n".join(block[j:])})
                    for n, buf in split_paras(block[j + 1:], para_regex, para_number):
                        out.append({"cite_id": "%s#protocol.%s" % (doc_id, n), "doc_id": doc_id,
                                    "kind": "paragraph", "label": "议定书第%s条" % n,
                                    "parent": doc_id + "#protocol", "text": "\n".join(buf)})
                    block = block[:j]
                    break
        cid = "%s#%s" % (doc_id, no)
        out.append({"cite_id": cid, "doc_id": doc_id, "kind": "article", "label": fmt_art % no, "parent": "",
                    "text": "\n".join(block)})
        for n, buf in split_paras(block[1:], para_regex, para_number):
            out.append({"cite_id": "%s.%s" % (cid, n), "doc_id": doc_id, "kind": "paragraph",
                        "label": fmt_para % (no, n), "parent": cid, "text": "\n".join(buf)})
    return out, mode


def slice_cn(doc_id, text):
    """Chinese legal text. Article mode (第X条) if present, else top-level item mode (一、二、)."""
    out, mode = slice_articles(doc_id, text, RE_CN_ART, cn2int, RE_CN_ITEM, cn2int,
                               "第%s条", "第%s条第%s款", "cn_article")
    if out:
        return out, mode
    for n, buf in split_paras(clean_lines(text), RE_CN_ITEM, cn2int):
        out.append({"cite_id": "%s#i%s" % (doc_id, n), "doc_id": doc_id, "kind": "item", "label": "第%s项" % n,
                    "parent": "", "text": "\n".join(buf)})
    return out, "cn_item" if out else "none"


def slice_en_treaty(doc_id, text):
    return slice_articles(doc_id, text, RE_EN_ART, int, RE_EN_PARA, int,
                          "Article %s", "Article %s(%s)", "en_article")


RE_HK_SEC = re.compile(r"^(\d+)([A-Z]{0,6})\.\t(\S.*)$")
RE_HK_SCH = re.compile(r"^(?:(Schedule|附表)\s*(\d+[A-Z]*)|(First|Second|Third|Fourth|Fifth|Sixth|Seventh|Eighth|Ninth|Tenth) Schedule)$")
ORDINAL = {"First": "1", "Second": "2", "Third": "3", "Fourth": "4", "Fifth": "5", "Sixth": "6", "Seventh": "7", "Eighth": "8", "Ninth": "9", "Tenth": "10"}


def hk_sch_no(line):
    m = RE_HK_SCH.match(line.strip())
    return (m.group(2) or ORDINAL[m.group(3)]) if m else None


def slice_hk_ordinance(doc_id, text):
    """e-Legislation RTF text: '14.<TAB>Charge of profits tax'. Section numbers never decrease in the body;
    numbered lines inside Schedules are not sections, so the body ends at the first Schedule heading."""
    lines = [ln.rstrip() for ln in text.splitlines() if ln.strip()]
    sch = [i for i, ln in enumerate(lines) if RE_HK_SCH.match(ln.strip())]
    body_end = sch[0] if sch else len(lines)
    heads, last = [], 0
    for i, ln in enumerate(lines[:body_end]):
        m = RE_HK_SEC.match(ln)
        if not m:
            continue
        n, title = int(m.group(1)), m.group(3).rstrip()
        if last <= n <= last + 3 and (title.endswith("etc.") or not title.endswith((".", ";", "—", ":", "。", "；"))):
            heads.append((i, m.group(1) + m.group(2)))
            last = n
    out, seen = [], set()
    for k, (i, no) in enumerate(heads):
        if no in seen:
            continue
        seen.add(no)
        stop = heads[k + 1][0] if k + 1 < len(heads) else body_end
        out.append({"cite_id": "%s#s%s" % (doc_id, no), "doc_id": doc_id, "kind": "section", "label": "s %s" % no,
                    "parent": "", "text": "\n".join(x.strip() for x in lines[i:stop])})
    seen_sch = set()
    for k, i in enumerate(sch):
        no = hk_sch_no(lines[i])
        if no in seen_sch:
            continue
        seen_sch.add(no)
        stop = sch[k + 1] if k + 1 < len(sch) else len(lines)
        out.append({"cite_id": "%s#sch%s" % (doc_id, no), "doc_id": doc_id, "kind": "schedule",
                    "label": "Schedule %s" % no, "parent": "", "text": "\n".join(x.strip() for x in lines[i:stop])})
    return out, "hk_section" if len(out) >= 20 else "none"


RE_NUM_PARA = re.compile(r"^(\d{1,3})\.(?:\s+\S.*)?$")
RE_DOT_PARA = re.compile(r"^(\d{1,2})\.(\d{1,2})(?:\s+\S.*)?$")
H_MARK = "@@H@@"


def best_run(lines, hits):
    """hits: [(line_index, key, restart)]. Split into runs at restarts; return the run with the most text."""
    runs, cur = [], []
    for h in hits:
        if h[2] and cur:
            runs.append(cur)
            cur = []
        cur.append(h)
    if cur:
        runs.append(cur)
    if not runs:
        return []

    def size(r):
        k = runs.index(r)
        end = runs[k + 1][0][0] if k + 1 < len(runs) else len(lines)
        return sum(len(x) for x in lines[r[0][0]:end])
    return max(runs, key=size)


def slice_numbered_paras(doc_id, text):
    """Guidance PDFs. Either '12. text' numbered consecutively (DIPN style) or '5.3 text' (e-Tax guide style)."""
    lines = clean_lines(text)
    hits, last = [], 0
    for i, ln in enumerate(lines):
        m = RE_NUM_PARA.match(ln)
        if not m:
            continue
        n = int(m.group(1))
        if n == 1 and last >= 3:
            hits.append((i, "1", True))
            last = 1
        elif n == last + 1 or (last == 0 and n == 2):
            hits.append((i, str(n), last == 0))
            last = n
    plain = best_run(lines, hits)
    hits, last = [], (0, 0)
    for i, ln in enumerate(lines):
        m = RE_DOT_PARA.match(ln)
        if not m:
            continue
        key = (int(m.group(1)), int(m.group(2)))
        if key == (1, 1) and last > (1, 1):
            hits.append((i, "1.1", True))
            last = key
        elif (key[0] == last[0] and key[1] == last[1] + 1) or (key[0] > last[0] and key[0] <= last[0] + 2 and key[1] == 1):
            hits.append((i, "%d.%d" % key, last == (0, 0)))
            last = key
    dotted = best_run(lines, hits)
    run = dotted if len(dotted) >= max(8, len(plain) // 2) else plain
    if len(run) < 8:
        return [], "none"
    out = []
    if run[0][0] > 0:
        out.append({"cite_id": doc_id + "#front", "doc_id": doc_id, "kind": "front", "label": "front matter",
                    "parent": "", "text": "\n".join(lines[:run[0][0]])})
    for k, (i, key, _) in enumerate(run):
        stop = run[k + 1][0] if k + 1 < len(run) else len(lines)
        out.append({"cite_id": "%s#para%s" % (doc_id, key), "doc_id": doc_id, "kind": "paragraph",
                    "label": "para %s" % key, "parent": "", "text": "\n".join(lines[i:stop])})
    return out, "numbered_para"


def slice_html_sections(doc_id, raw):
    """Web pages: one clause per heading (h1-h4) inside the main container."""
    soup = BeautifulSoup(raw, "lxml")
    for t in soup(["script", "style", "noscript", "header", "footer", "nav", "form", "iframe"]):
        t.decompose()
    body = soup.body or soup

    def flat(node):
        return " ".join(node.get_text(" ").split())
    faq = body.find_all(class_=re.compile("question_container"))
    if len(faq) >= 3:
        out = []
        for n, q in enumerate(faq, 1):
            t = q.find(class_=re.compile("faq_title"))
            body_lines = [" ".join(x.split()) for x in q.get_text("\n").splitlines() if x.strip()]
            out.append({"cite_id": "%s#q%d" % (doc_id, n), "doc_id": doc_id, "kind": "qa",
                        "label": (flat(t) if t else "Q%d" % n)[:80], "parent": "", "text": "\n".join(body_lines)})
        return out, "html_faq"
    for h in body.find_all(["h1", "h2", "h3", "h4"]):
        h.insert_before("\n%s " % H_MARK)
    for para in body.find_all("p"):
        st = para.find("strong")
        if st and 3 < len(flat(para)) < 120 and flat(para) == flat(st):
            para.insert_before("\n%s " % H_MARK)
    lines = [re.sub(r"[ \t 　]+", " ", ln).strip() for ln in body.get_text("\n").splitlines()]
    lines = [ln for ln in lines if ln]
    sections, title, buf = [], None, []
    i = 0
    while i < len(lines):
        if lines[i].startswith(H_MARK):
            if title is not None and buf:
                sections.append((title, buf))
            rest = lines[i][len(H_MARK):].strip()
            if not rest and i + 1 < len(lines):
                i += 1
                rest = lines[i]
            title, buf = rest, [rest]
        elif title is not None:
            buf.append(lines[i])
        i += 1
    if title is not None and buf:
        sections.append((title, buf))
    sections = [(t, b) for t, b in sections if sum(len(x) for x in b) > len(t) + 40]   # drop empty menu headings
    if len(sections) < 3:
        return [], "none"
    out = []
    for n, (t, b) in enumerate(sections, 1):
        out.append({"cite_id": "%s#h%d" % (doc_id, n), "doc_id": doc_id, "kind": "section", "label": t[:80],
                    "parent": "", "text": "\n".join(b)})
    return out, "html_section"


SSO_ANCHOR = re.compile(r"^(pr[0-9]+[A-Z]*-|Sc[0-9]+[A-Z]*-)$")
SSO_REPEALED = re.compile(r"^\d+[A-Z]*\.\s*\[\s*Repealed by [^\]]*\]\s*$", re.S)
SSO_DROPPED = []                     # (cite_id, text): repealed-provision stubs left out of the clause layer (SSO, e-Legislation)
HK_REPEALED = re.compile(r"^(?:\d+[A-Z]*\.|Schedule \d+[A-Z]*)\s*\((?:Repealed|Omitted)[^)]*\)\s*$|"
                         r"^(?:\d+[A-Z]*\.|附表\d+[A-Z]*)\s*\(由[^)]*(?:廢除|刪除|删除)\)\s*$")


def sso_segments(soup):
    """Text from each SSO anchor to the next one, in document order: the fallback when a provision has no box of
    its own (the nearest enclosing div would otherwise be the whole Act)."""
    from bs4 import NavigableString, Tag
    seg, cur = {}, None
    for el in soup.descendants:
        if isinstance(el, Tag):
            if el.get("id") and SSO_ANCHOR.match(el["id"]):
                cur = el["id"]
                seg.setdefault(cur, [])
        elif isinstance(el, NavigableString) and cur and el.parent.name not in ("script", "style"):
            seg[cur].append(str(el))
    return {k: "\n".join(x.strip() for x in "\n".join(v).splitlines() if x.strip()) for k, v in seg.items()}


def slice_sso(doc_id, raw):
    """Singapore Statutes Online HTML: one clause per section, anchored by id='pr<no>-'. A provision's text is its
    own prov1 box when that box holds no other anchor, else the anchor-to-anchor segment. Repealed-section stubs
    ('18A. [Repealed by Act 21 of 2003]') are not current law and stay out of the clause layer."""
    soup = BeautifulSoup(raw, "lxml")
    segments = None
    out, seen = [], set()
    for pattern, key, kind, label, box_class in ((r"^pr[0-9]+[A-Z]*-$", "s", "section", "s %s", "prov1"),
                                                 (r"^Sc[0-9]+[A-Z]*-$", "sch", "schedule", "Schedule %s", None)):
        for hdr in soup.find_all(id=re.compile(pattern)):
            no = hdr["id"][2:-1]
            if (key, no) in seen:
                continue
            seen.add((key, no))
            box = hdr.find_parent("div", class_=re.compile(box_class)) if box_class else hdr.find_parent("div")
            if box is None or len(box.find_all(id=SSO_ANCHOR)) > 1:
                if segments is None:
                    segments = sso_segments(soup)
                text = segments.get(hdr["id"], "")
            else:
                text = "\n".join(s.strip() for s in box.get_text("\n").splitlines() if s.strip())
            cite = "%s#%s%s" % (doc_id, key, no)
            if SSO_REPEALED.match(text.replace("\n", " ")):
                SSO_DROPPED.append((cite, text.replace("\n", " ")))
                continue
            out.append({"cite_id": cite, "doc_id": doc_id, "kind": kind, "label": label % no, "parent": "", "text": text})
    return out, "sso_section" if out else "none"


def slice_xfa(doc_id, text):
    """One clause per top-level subform of a dynamic form, e.g. #Page1_E; header, footer and buttons are skipped."""
    groups, order = {}, []
    for ln in text.splitlines():
        path = ln.split(" | ", 1)[0]
        parts = path.split("/")
        key = parts[1] if len(parts) > 1 else "root"
        if key.lower().startswith(("subform_header", "subform_footer", "landingpage")) or key == "root":
            continue
        if key not in groups:
            groups[key] = []
            order.append(key)
        groups[key].append(ln)
    out = [{"cite_id": "%s#%s" % (doc_id, k), "doc_id": doc_id, "kind": "subform", "label": k, "parent": "",
            "text": "\n".join(groups[k])} for k in order]
    return out, "xfa_subform" if out else "none"


def slice_pages(doc_id, pages):
    return [{"cite_id": "%s#p%d" % (doc_id, i + 1), "doc_id": doc_id, "kind": "page", "label": "p.%d" % (i + 1),
             "parent": "", "text": t.strip()} for i, t in enumerate(pages) if t.strip()], "page"


def slice_doc(row, text, pages, all_ids):
    doc_id, typ = row["id"], row["type"]
    if ".frag" in doc_id or (doc_id + ".full") in all_ids:
        return [], "skipped_part"                      # raw piece of a document that is sliced as <id>.full
    if row.get("alias_of"):
        return [], "alias"                             # same bytes as another id
    if row.get("text_status") == "xfa_template":
        return slice_xfa(doc_id, text)
    if doc_id.startswith(("sg.ita", "sg.ia", "sg.ha", "sg.sda", "sg.gsta", "sg.memta")) and row["ext"] == "html":   # Singapore Statutes Online Acts
        out, mode = slice_sso(doc_id, (ROOT / row["file"]).read_bytes())
        if out:
            return out, mode
    if doc_id.startswith("hk.cap"):                                                   # e-Legislation chapters
        out, mode = slice_hk_ordinance(doc_id, text)
        if out and mode != "none":
            return out, mode
    scope = text
    if pages and doc_id in PAGE_RANGE:
        a, b = PAGE_RANGE[doc_id]
        scope = "\n".join(pages[a - 1:b])
    sample = scope[:20000]
    cn = len(re.findall(r"[\u4e00-\u9fff]", sample)) > 0.2 * max(1, len(sample))
    is_treaty = doc_id.startswith("treaty.")
    attached_text = typ == "attachment" and re.search(r"办法|规定|指南|细则|条例|准则|公约|协定", row["title"] or "") and "表" not in (row["title"] or "")
    legal = typ in ("law", "regulation", "guidance") or is_treaty or bool(attached_text)   # an attached 办法 slices like its parent
    if cn and legal:
        out, mode = slice_cn(doc_id, scope)
        if len(out) >= 3:
            return out, mode
    if not cn and is_treaty:
        out, mode = slice_en_treaty(doc_id, scope)
        if len(out) >= 3:
            return out, mode
    if pages and typ in ("guidance", "guide", "example", "page") and not cn:
        out, mode = slice_numbered_paras(doc_id, text)
        if out:
            return out, mode
    if pages:
        return slice_pages(doc_id, pages)
    if row.get("real_kind") in ("html", "htm"):
        out, mode = slice_html_sections(doc_id, (ROOT / row["file"]).read_bytes())
        if out:
            return out, mode
    return [{"cite_id": doc_id + "#all", "doc_id": doc_id, "kind": "whole", "label": "全文", "parent": "",
             "text": text}], "whole"


# ---------------------------------------------------------------- 4. params check
def norm(s):
    return re.sub(r"\s+", "", s or "")


def drop_superseded(clauses):
    """superseded.csv: doc_id, clause ('13.4', or '24*' for a prefix), by, since. Those clauses are not kept."""
    rules = read_csv(ROOT / "superseded.csv")
    if not rules:
        return clauses, 0
    cut = {}                                           # "i8@2" or "13.5@4-": remove those lines of the clause
    for r in rules:
        if "@" in r["clause"]:
            name, span = r["clause"].split("@")
            a, _, b = span.partition("-")
            cut[r["doc_id"] + "#" + name] = (int(a), int(b) if b else (10 ** 6 if "-" in span else int(a)))
    phrases = [r for r in rules if r.get("phrase")]     # remove an exact phrase from one clause
    rules = [r for r in rules if "@" not in r["clause"] and not r.get("phrase")]
    exact = {r["doc_id"] + "#" + r["clause"] for r in rules if not r["clause"].endswith("*")}
    prefix = [r["doc_id"] + "#" + r["clause"][:-1] for r in rules if r["clause"].endswith("*")]
    kept = [c for c in clauses if c["cite_id"] not in exact and not any(c["cite_id"].startswith(x) for x in prefix)]
    n_cut = 0
    for c in kept:
        if c["cite_id"] in cut:
            a, b = cut[c["cite_id"]]
            lines = c["text"].splitlines()
            c["text"] = "\n".join(ln for i, ln in enumerate(lines, 1) if not a <= i <= b)
            c["chars"] = len(c["text"])
            n_cut += 1
    by_id = {c["cite_id"]: c for c in kept}
    for r in phrases:
        c = by_id.get(r["doc_id"] + "#" + r["clause"])
        targets = [c] if c is not None else []
        if c is None and r["clause"] == "all":            # the document is sliced now: cut the phrase where it sits
            targets = [x for x in kept if x["doc_id"] == r["doc_id"] and r["phrase"] in x["text"]]
        if not targets or any(r["phrase"] not in x["text"] for x in targets):
            raise SystemExit("superseded.csv: phrase not found in %s#%s: %s" % (r["doc_id"], r["clause"], r["phrase"][:30]))
        for x in targets:
            x["text"] = "\n".join(ln for ln in x["text"].replace(r["phrase"], "").splitlines() if ln.strip())
            x["chars"] = len(x["text"])
            n_cut += 1
    return kept, len(clauses) - len(kept) + n_cut


def validity_label(row):
    """Status shown on fgk.chinatax.gov.cn pages: <p class="arc_date"><span class="xg">全文有效</span>."""
    if row["jur"] != "cn" or row.get("real_kind") not in ("html", "htm"):
        return ""
    m = re.search(r'class="xg">([^<]*)<', (ROOT / row["file"]).read_bytes().decode("utf-8", "replace"))
    return m.group(1).strip() if m else ""


def check_params(params, clause_text):
    issues = []
    for p in params:
        cid = p.get("cite_id", "")
        if cid not in clause_text:
            issues.append("param %s: cite_id not found: %s" % (p.get("param_id"), cid))
        elif p.get("quote") and norm(p["quote"]) not in norm(clause_text[cid]):
            issues.append("param %s: quote not found verbatim in %s" % (p.get("param_id"), cid))
    return issues


def check_form_fields(rows, doc_ids):
    """form_fields.csv: every label must occur verbatim in the form's text; the form must still be in the library."""
    issues, cache = [], {}
    for r in rows:
        fid = r["form_id"]
        if fid not in doc_ids:
            issues.append("form_fields: form not in library: %s" % fid)
            continue
        if fid not in cache:
            f = TEXT / (fid + ".txt")
            cache[fid] = norm(f.read_text(encoding="utf-8", errors="replace")) if f.exists() else ""
        if norm(r["label"]) not in cache[fid]:
            issues.append("form_fields: label not found verbatim in %s: %s" % (fid, r["label"][:40]))
    return issues


# ---------------------------------------------------------------- main
def main():
    docs, problems = merge_manifests()
    TEXT.mkdir(exist_ok=True)
    all_ids = {r["id"] for r in docs}
    kinds = {r["id"]: real_kind(ROOT / r["file"], r["ext"]) for r in docs}
    soffice_convert([(ROOT / r["file"], kinds[r["id"]], r["sha256"][:16]) for r in docs
                     if kinds[r["id"]] in ("doc", "rtf", "xls", "ppt", "pptx")])
    clauses = []
    currency = {c["id"]: (c["currency"], c["note"]) for c in read_csv(ROOT / "currency.csv")}
    partly_superseded = {x["doc_id"] for x in read_csv(ROOT / "superseded.csv")}
    for r in docs:
        text, pages, status = to_text(r, kinds[r["id"]])
        r["real_kind"], r["text_status"], r["text_chars"] = kinds[r["id"]], status, len(text)
        r["validity_label"] = validity_label(r)
        r["currency"], r["currency_note"] = currency.get(r["id"], ("", ""))
        tfile = TEXT / (r["id"] + ".txt")
        if text and r["id"] not in partly_superseded:
            tfile.write_text(text, encoding="utf-8")
        elif tfile.exists():
            tfile.unlink()                             # the clause layer is the only derived copy for these
        cl, mode = slice_doc(r, text, pages, all_ids) if text else ([], "none")
        r["slice_mode"], r["clauses"] = mode, len(cl)
        clauses.extend(cl)
    kept = []
    for c in clauses:                                  # repealed-provision placeholders are not current law
        head = re.split(r"_{5,}", c["text"])[0].replace("\n", " ").replace("\t", " ").strip()
        if HK_REPEALED.match(head) and c["doc_id"].startswith("hk."):
            SSO_DROPPED.append((c["cite_id"], head))
        else:
            kept.append(c)
    clauses = kept
    seen, uniq = set(), []
    for c in clauses:
        if c["cite_id"] in seen:
            problems.append("duplicate cite_id dropped: %s" % c["cite_id"])
            continue
        seen.add(c["cite_id"])
        c["chars"] = len(c["text"])
        uniq.append(c)
    clauses, n_superseded = drop_superseded(uniq)

    doc_fields = list(docs[0].keys()) if docs else []
    write_csv(ROOT / "manifest.csv", doc_fields, docs)
    clause_fields = ["cite_id", "doc_id", "kind", "label", "parent", "chars", "text"]
    write_csv(ROOT / "clauses.csv", clause_fields, clauses)

    params = read_csv(ROOT / "params.csv")
    issues = check_params(params, {c["cite_id"]: c["text"] for c in clauses})
    form_fields = read_csv(ROOT / "form_fields.csv")
    form_fields = [r for r in form_fields if r["form_id"] in all_ids] if form_fields else []
    issues += check_form_fields(form_fields, all_ids)
    gaps = read_csv(ROOT / "gaps.csv")                 # dictionary gaps checked against the text: same quote rule as params
    issues += [x.replace("param ", "gap ") for x in
               check_params([dict(g, param_id=g["new_id"]) for g in gaps], {c["cite_id"]: c["text"] for c in clauses})]
    calendar = []
    for cal in sorted(list((ROOT / "calendars").glob("*.csv")) + list((ROOT / "labels").glob("*.csv"))):
        page_text = {}
        for r in read_csv(cal):
            if "doc_id" not in r or "quote" not in r:      # labels taken by hand from a tool have no page to quote
                continue
            d = r["doc_id"]
            if d not in page_text:                     # whole page, not only the main container
                src = next((ROOT / x["file"] for x in docs if x["id"] == d), None)
                page_text[d] = norm(BeautifulSoup(src.read_bytes(), "html.parser").get_text(" ")) if src else ""   # lxml truncates some gov.cn pages
            if norm(r["quote"]) not in page_text[d]:
                issues.append("%s: quote not found in %s: %s" % (cal.name, d, r["quote"][:40]))
            if cal.parent.name == "calendars":
                calendar.append(dict(r, calendar=cal.stem))

    db = ROOT / "official.sqlite"
    if db.exists():
        db.unlink()
    con = sqlite3.connect(str(db))
    for name, fields, rows in (("doc", doc_fields, docs), ("clause", clause_fields, clauses),
                               ("param", list(params[0].keys()) if params else [], params),
                               ("form_field", list(form_fields[0].keys()) if form_fields else [], form_fields),
                               ("gap", list(gaps[0].keys()) if gaps else [], gaps),
                               ("calendar", list(calendar[0].keys()) if calendar else [], calendar)):
        if not fields:
            continue
        con.execute("CREATE TABLE %s (%s)" % (name, ", ".join('"%s" TEXT' % f for f in fields)))
        con.executemany("INSERT INTO %s VALUES (%s)" % (name, ",".join("?" * len(fields))),
                        [[str(r.get(f, "")) for f in fields] for r in rows])
    con.commit()
    con.close()

    blocked = []
    for p in sorted(ROOT.glob("blocked_*.csv")):
        blocked.extend(read_csv(p))
    blocked = [b for b in blocked if b["id"] not in all_ids]
    listed = [r for r in docs if ".frag" not in r["id"]]

    lines = ["# STATUS", "", "由 `tools/build.py` 生成，不要手改。", "",
             "| 项 | 数量 |", "|---|---|",
             "| 入库文件 | %d |" % len(docs),
             "| 其中分片原件 `.fragNN` | %d |" % (len(docs) - len(listed)),
             "| 官方域名 | %d |" % sum(1 for r in docs if r.get("official_domain") == "yes"),
             "| 本地拼装或非原件 | %d |" % sum(1 for r in docs if r.get("verbatim") != "raw"),
             "| 已转文本 | %d |" % sum(1 for r in docs if r["text_status"] == "ok"),
             "| 条款片 | %d |" % len(clauses),
             "| 因被取代或废止而不保留的条款片 | %d |" % n_superseded,
             "| 参数 | %d |" % len(params),
             "| 表单栏位映射 | %d |" % len(form_fields),
             "| 字典缺口核证 | %d |" % len(gaps),
             "| 日历行 | %d |" % len(calendar),
             "| 未取到 | %d |" % len(blocked), "",
             "## 入库文件", "", "| id | type | kind | bytes | text | slice | clauses | via | title |",
             "|---|---|---|---|---|---|---|---|---|"]
    for r in sorted(listed, key=lambda x: x["id"]):
        lines.append("| %s | %s | %s | %s | %s | %s | %s | %s | %s |" % (
            r["id"], r["type"], r["real_kind"], r["bytes"], r["text_status"], r["slice_mode"], r["clauses"],
            r["via"], r["title"].replace("|", "/")[:60]))
    lines += ["", "## 未取到", ""]
    if blocked:
        lines += ["| id | title | reason | alt_tried |", "|---|---|---|---|"]
        for b in blocked:
            lines.append("| %s | %s | %s | %s |" % (b["id"], b["title"].replace("|", "/")[:50],
                                                   b["reason"].replace("|", "/")[:160],
                                                   b["alt_tried"].replace("|", "/")[:160]))
    else:
        lines.append("无。")
    if SSO_DROPPED:
        lines += ["", "## 未入条款层的已废止条文（法规在线文本中的废止占位）", ""]
        lines += ["- `%s`：%s" % (c, t[:80]) for c, t in SSO_DROPPED]
    lines += ["", "## 检查", ""]
    lines += ["- " + x for x in (problems + issues)] or ["- 无问题。"]
    (ROOT / "STATUS.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    print("docs=%d clauses=%d params=%d blocked=%d problems=%d param_issues=%d" % (
        len(docs), len(clauses), len(params), len(blocked), len(problems), len(issues)))
    for x in problems + issues:
        print("  !", x)
    write_csv(ROOT / "repealed_stubs.csv", ["cite_id", "text"], [{"cite_id": c, "text": t} for c, t in SSO_DROPPED])
    import graph                                       # knowledge layer: identities, units, relations, anchors, index
    graph.main(sys.argv[1:])


if __name__ == "__main__":
    main()
